"""Vuelca las celdas crudas de las planillas reales, hoja por hoja.

Sirve para verificar la lectura (Lectura.gs) sin Apps Script: el harness de
Node parsea estas mismas celdas y compara las respuestas que obtiene contra
las que obtuvo el motor Python (guardadas en referencia-python.json).

Uso:  python3 dump-celdas.py       # requiere haber corrido dump-referencia.py
"""
import json
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
PLANILLAS = RAIZ / "compartido" / "instrumentos"
SALIDA = Path(__file__).parent / "celdas-python.json"

HOJAS = ["NEO", "CELID-A", "POTENLID", "CAMIN-A", "CONLID-A"]


def celda(valor):
    """Google Sheets entrega '' donde openpyxl entrega None."""
    if valor is None:
        return ""
    if isinstance(valor, (int, float, str, bool)):
        return valor
    return str(valor)


def grilla(ws):
    filas = [[celda(v) for v in fila] for fila in ws.iter_rows(values_only=True)]
    # Las hojas traen cientos de filas vacías al final; no aportan nada.
    while filas and not any(v != "" for v in filas[-1]):
        filas.pop()
    return filas


def main():
    rutas = sorted(PLANILLAS.glob("Planilla de Preguntas*.xlsx"))
    volcados = []
    for ruta in rutas:
        wb = openpyxl.load_workbook(ruta)
        faltan = [h for h in HOJAS if h not in wb.sheetnames]
        if faltan:
            print(f"  ✗ {ruta.name}: le faltan las hojas {faltan}")
            continue
        volcados.append({
            "planilla": ruta.name,
            "grillas": {h: grilla(wb[h]) for h in HOJAS},
        })
        print(f"  ✓ {ruta.name}")

    SALIDA.write_text(json.dumps(volcados, ensure_ascii=False), encoding="utf-8")
    print(f"\n{len(volcados)} planilla(s) → {SALIDA.relative_to(RAIZ)}")
    return 0 if volcados else 1


if __name__ == "__main__":
    sys.exit(main())
