"""Genera el JSON de referencia contra el que se verifica el port a JS.

No reimplementa la corrección: extrae el bloque de código real de
run_engine.py (desde la lectura de la planilla hasta el vector del radar)
y lo ejecuta tal cual. Así la referencia es el motor en producción, no una
copia que podría divergir sin que nadie lo note.

Las planillas reales recorren pocas celdas de los baremos, así que además se
generan casos sintéticos con respuestas al azar: son los que recorren los
bordes de las tablas, donde un dígito mal copiado pasaría inadvertido.

Uso:  python3 dump-referencia.py                # planillas reales + 500 sintéticos
      python3 dump-referencia.py --fuzz 5000    # más casos sintéticos
      python3 dump-referencia.py <xlsx>...      # planillas puntuales, sin fuzz
"""
import json
import random
import re
import sys
import textwrap
import types
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
MOTOR = RAIZ / "01-legacy-python" / "run_engine.py"
PLANILLAS = RAIZ / "compartido" / "instrumentos"
SALIDA = Path(__file__).parent / "referencia-python.json"

INICIO = "wb = openpyxl.load_workbook(xlsx_path)"
FIN = "ideal = [90, 85, 90, 90, 85, 75, 70, 75, 80, 90, 85, 75, 85, 85]"


def bloque_de_correccion():
    """Recorta el fragmento de run_engine.py que corrige y arma el radar."""
    lineas = MOTOR.read_text(encoding="utf-8").splitlines()
    desde = next(i for i, l in enumerate(lineas) if INICIO in l)
    hasta = next(i for i, l in enumerate(lineas) if FIN in l)
    return textwrap.dedent("\n".join(lineas[desde:hasta + 1]))


def corregir(codigo, lector, xlsx_path=""):
    ambito = {
        "openpyxl": lector,
        "re": re,
        "xlsx_path": str(xlsx_path),
        "pc": lambda *_: None,
    }
    exec(codigo, ambito)
    return ambito


# ── Planillas sintéticas ────────────────────────────────────────────
# Un libro mínimo con la misma forma que espera run_engine, para poder
# alimentar el motor real con respuestas al azar sin generar archivos.

ESCALAS = {"celid": (1, 5), "potenlid": (1, 5), "camin": (1, 7), "conlid": (1, 5)}
CANTIDAD_ITEMS = {"neo": 60, "celid": 34, "potenlid": 9, "camin": 12, "conlid": 18}
VACIA = (None, None, None, None)


class _Hoja:
    def __init__(self, filas):
        self._filas = filas

    def iter_rows(self, min_row=1, values_only=False):
        return iter(self._filas[min_row - 1:])


class _Libro:
    def __init__(self, hojas):
        self._hojas = hojas

    def __getitem__(self, nombre):
        return self._hojas[nombre]


def _libro_sintetico(respuestas):
    """Replica la disposición de filas que lee run_engine en cada hoja."""
    neo = [VACIA] * 5 + [(f"{i}. ítem", respuestas["neo"][i], None, None) for i in range(1, 61)]
    celid = [VACIA] * 3 + [(i, None, None, respuestas["celid"][i]) for i in range(1, 35)]
    resto = {}
    for hoja, clave in (("POTENLID", "potenlid"), ("CAMIN-A", "camin"), ("CONLID-A", "conlid")):
        filas = [VACIA] * 2
        filas += [(i, None, respuestas[clave][i], None) for i in range(1, CANTIDAD_ITEMS[clave] + 1)]
        resto[hoja] = _Hoja(filas)
    return _Libro({"NEO": _Hoja(neo), "CELID-A": _Hoja(celid), **resto})


SESGOS = ("uniforme", "bajo", "alto", "extremo")


def _valor(rnd, bajo, alto, sesgo):
    """Un sesgo por instrumento: con respuestas uniformes los puntajes se
    agolpan en el medio y los cortes P5/P95 de los baremos nunca se visitan."""
    if sesgo == "bajo":
        return rnd.randint(bajo, min(bajo + 1, alto))
    if sesgo == "alto":
        return rnd.randint(max(alto - 1, bajo), alto)
    if sesgo == "extremo":
        return rnd.choice((bajo, alto))
    return rnd.randint(bajo, alto)


def respuestas_borde(letra_neo, extremo):
    """Todo el instrumento en su valor mínimo o máximo: fuerza las filas
    extremas de los baremos, que el azar casi nunca alcanza."""
    respuestas = {"neo": {i: letra_neo for i in range(1, 61)}}
    for clave, (bajo, alto) in ESCALAS.items():
        valor = alto if extremo == "max" else bajo
        respuestas[clave] = {i: valor for i in range(1, CANTIDAD_ITEMS[clave] + 1)}
    return respuestas


def respuestas_al_azar(rnd):
    sesgo_neo = rnd.choice(SESGOS)
    letras = {"bajo": "AB", "alto": "DE", "extremo": "AE"}.get(sesgo_neo, "ABCDE")
    respuestas = {"neo": {i: rnd.choice(letras) for i in range(1, 61)}}
    for clave, (bajo, alto) in ESCALAS.items():
        sesgo = rnd.choice(SESGOS)
        respuestas[clave] = {
            i: _valor(rnd, bajo, alto, sesgo) for i in range(1, CANTIDAD_ITEMS[clave] + 1)
        }
    return respuestas


def a_claves_str(d):
    """JSON exige claves string; el harness JS las vuelve a indexar por número."""
    return {str(k): v for k, v in d.items()}


def armar_caso(nombre, v):
    return {
            "planilla": nombre,
            "respuestas": {
                "neo": a_claves_str(v["neo_resp"]),
                "celid": a_claves_str(v["celid"]),
                "potenlid": a_claves_str(v["poten"]),
                "camin": a_claves_str(v["cam"]),
                "conlid": a_claves_str(v["con"]),
            },
            "esperado": {
                "neo": {"raw": v["neo_raw"], "t": v["neo_t"], "nivel": v["neo_niv"]},
                "celid": {
                    "valor": {
                        "Carisma": v["carisma"], "EstimInt": v["est_int"],
                        "Inspir": v["inspir"], "ConsInd": v["cons_ind"],
                        "TransfTot": v["transf_total"], "RecCont": v["rec_cont"],
                        "DirExc": v["dir_exc"], "TransTot": v["trans_total"],
                        "Laissez": v["laissez"],
                    },
                    "percentil": v["cel_p"],
                },
                "potenlid": {
                    "valor": {"Intr": v["m_intr"], "Extr": v["m_extr"], "Soc": v["m_soc"]},
                    "percentil": v["pot_p"],
                },
                "camin": {
                    "valor": {
                        "Dir": v["directivo"], "Cons": v["considerado"],
                        "Part": v["participativo"], "Or": v["orientado"],
                    },
                    "percentil": v["cam_p"],
                },
                "conlid": {
                    "valor": {"Tar": v["tarea"], "Rel": v["rel"], "Camb": v["camb"]},
                    "percentil": v["con_p"],
                },
                "radar": {"evaluado": v["evaluado_vals"], "ideal": v["ideal"]},
            },
    }


def main():
    argv = sys.argv[1:]
    fuzz = 500
    if "--fuzz" in argv:
        i = argv.index("--fuzz")
        fuzz = int(argv[i + 1])
        argv = argv[:i] + argv[i + 2:]
    rutas = [Path(a) for a in argv]
    if not rutas:
        rutas = sorted(PLANILLAS.glob("Planilla de Preguntas*.xlsx"))
    else:
        fuzz = 0

    codigo = bloque_de_correccion()
    casos = []

    for ruta in rutas:
        try:
            v = corregir(codigo, openpyxl, ruta)
        except Exception as e:
            print(f"  ✗ {ruta.name}: {type(e).__name__}: {e}")
            continue
        casos.append(armar_caso(ruta.name, v))
        print(f"  ✓ {ruta.name}")

    def agregar(nombre, respuestas):
        lector = types.SimpleNamespace(load_workbook=lambda _p, libro=_libro_sintetico(respuestas): libro)
        casos.append(armar_caso(nombre, corregir(codigo, lector)))

    if fuzz:
        for letra in ("A", "E"):
            for extremo in ("min", "max"):
                agregar(f"borde-neo{letra}-{extremo}", respuestas_borde(letra, extremo))
        print("  ✓ 4 planillas de borde (todo mínimo / todo máximo)")

        rnd = random.Random(20260723)  # semilla fija: los fallos son reproducibles
        for n in range(fuzz):
            agregar(f"sintetica-{n + 1:04d}", respuestas_al_azar(rnd))
        print(f"  ✓ {fuzz} planillas sintéticas")

    SALIDA.write_text(json.dumps(casos, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n{len(casos)} caso(s) → {SALIDA.relative_to(RAIZ)}")
    return 0 if casos else 1


if __name__ == "__main__":
    sys.exit(main())
