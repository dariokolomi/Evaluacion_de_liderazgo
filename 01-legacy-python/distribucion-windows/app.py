"""Servidor web CCHH — Generador de Informes de Liderazgo."""
from flask import Flask, render_template, jsonify, request, Response, send_file
import os, json, time, threading, uuid
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict

app = Flask(__name__)
BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
HISTORY    = os.path.join(BASE_DIR, '.runs_history.json')
QUALITY_XL = os.path.join(BASE_DIR, 'CALIDAD DE INFORMES.xlsx')

# ── Estado en memoria de corridas activas ──
_runs = {}   # run_id -> {'events': [], 'done': bool}
_lock = threading.Lock()


# ════════════════════════════════════════
# Utilidades de persistencia
# ════════════════════════════════════════

def load_history():
    try:
        with open(HISTORY, encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return []

def save_history(h):
    with open(HISTORY, 'w', encoding='utf-8') as f:
        json.dump(h, f, indent=2, ensure_ascii=False)

def rebuild_excel():
    history = load_history()
    rated   = [r for r in history if r.get('rating')]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Calificaciones'
    headers = ['Fecha', 'Evaluado', 'Modelo', 'Planilla', 'Informe Generado', 'Tiempo (s)', 'Calificación', 'Comentario']
    hfill = PatternFill('solid', fgColor='2E5496')
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.font      = Font(bold=True, color='FFFFFF')
        cell.fill      = hfill
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 22
    for row_num, rec in enumerate(rated, 2):
        ws.cell(row_num, 1, rec['timestamp'][:19].replace('T', ' '))
        ws.cell(row_num, 2, rec.get('evaluado', ''))
        ws.cell(row_num, 3, rec.get('modelo', ''))
        ws.cell(row_num, 4, rec.get('xlsx', ''))
        ws.cell(row_num, 5, rec.get('output', ''))
        ws.cell(row_num, 6, rec.get('elapsed', 0))
        rating = rec.get('rating', 0)
        ws.cell(row_num, 7, rating)
        ws.cell(row_num, 8, rec.get('comment', ''))
        color = 'E2EFDA' if rating >= 4 else ('FFEB9C' if rating == 3 else 'FCE4D6')
        ws.cell(row_num, 7).fill = PatternFill('solid', fgColor=color)
    wb.save(QUALITY_XL)


# ════════════════════════════════════════
# Rutas principales
# ════════════════════════════════════════

@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/files')
def list_files():
    all_f = os.listdir(BASE_DIR)
    EXCLUIR_XLSX = {'calidad de informes.xlsx'}
    modelos   = sorted([f for f in all_f if f.lower().endswith('.docx') and not f.lower().startswith('informe')])
    planillas = sorted([f for f in all_f if f.lower().endswith('.xlsx') and f.lower() not in EXCLUIR_XLSX])
    resp = jsonify({'modelos': modelos, 'planillas': planillas})
    resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    resp.headers['Pragma'] = 'no-cache'
    return resp


@app.route('/api/last-run')
def last_run():
    h = load_history()
    return jsonify(h[-1] if h else None)


@app.route('/api/validate', methods=['POST'])
def validate():
    xlsx = request.json.get('xlsx', '')
    path = os.path.join(BASE_DIR, xlsx)
    if not os.path.exists(path):
        return jsonify({'error': 'Archivo no encontrado'}), 404
    required = ['NEO', 'CELID-A', 'POTENLID', 'CAMIN-A', 'CONLID-A']
    results  = {}
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in required:
            if sheet not in wb.sheetnames:
                results[sheet] = {'ok': False, 'rows': 0, 'msg': 'Hoja no encontrada'}
            else:
                ws   = wb[sheet]
                rows = sum(1 for row in ws.iter_rows(min_row=3, values_only=True)
                           if any(v is not None for v in row))
                results[sheet] = {'ok': rows > 5, 'rows': rows,
                                  'msg': f'{rows} filas con datos' if rows > 5 else 'Datos insuficientes'}
        wb.close()
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    return jsonify(results)


@app.route('/api/estimated-time')
def estimated_time():
    times = [r['elapsed'] for r in load_history() if r.get('elapsed', 0) > 0]
    if times:
        return jsonify({'estimated': round(sum(times) / len(times), 1), 'samples': len(times)})
    return jsonify({'estimated': None, 'samples': 0})


@app.route('/api/quality-by-model')
def quality_by_model():
    rated    = [r for r in load_history() if r.get('rating')]
    by_model = defaultdict(list)
    for r in rated:
        by_model[r['modelo']].append(r['rating'])
    return jsonify({m: round(sum(v)/len(v), 2) for m, v in by_model.items()})


# ════════════════════════════════════════
# Ejecución
# ════════════════════════════════════════

@app.route('/api/run', methods=['POST'])
def run_report():
    data     = request.json
    xlsx     = data.get('xlsx', '').strip()
    modelo   = data.get('modelo', '').strip()
    evaluado = (data.get('evaluado', '') or 'Evaluado').strip()
    if not xlsx or not modelo:
        return jsonify({'error': 'Faltan parámetros'}), 400
    for f in [xlsx, modelo]:
        if not os.path.exists(os.path.join(BASE_DIR, f)):
            return jsonify({'error': f'No se encontró: {f}'}), 404

    run_id      = uuid.uuid4().hex[:8]
    ts          = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_name = f"INFORME_{evaluado.replace(' ', '_')}_{ts}.docx"
    radar_name  = f".radar_{run_id}.png"

    with _lock:
        _runs[run_id] = {'events': [], 'done': False}

    t = threading.Thread(
        target=_bg_run,
        args=(run_id, xlsx, modelo, evaluado, output_name, radar_name),
        daemon=True
    )
    t.start()
    return jsonify({'run_id': run_id})


def _emit(run_id, pct, msg, status='running'):
    with _lock:
        if run_id in _runs:
            _runs[run_id]['events'].append({'pct': pct, 'msg': msg, 'status': status})
            if status in ('done', 'error'):
                _runs[run_id]['done'] = True


def _bg_run(run_id, xlsx, modelo, evaluado, output_name, radar_name):
    from run_engine import run_informe
    start = time.time()
    try:
        run_informe(
            os.path.join(BASE_DIR, xlsx),
            os.path.join(BASE_DIR, modelo),
            os.path.join(BASE_DIR, output_name),
            os.path.join(BASE_DIR, radar_name),
            evaluado,
            lambda pct, msg: _emit(run_id, pct, msg)
        )
        elapsed = round(time.time() - start, 1)
        history = load_history()
        history.append({
            'run_id':    run_id,
            'timestamp': datetime.now().isoformat(),
            'xlsx':      xlsx,
            'modelo':    modelo,
            'evaluado':  evaluado,
            'output':    output_name,
            'elapsed':   elapsed,
            'rating':    None,
            'comment':   ''
        })
        save_history(history)
        _emit(run_id, 100, f'Informe listo en {elapsed}s — {output_name}', 'done')
    except Exception as e:
        import traceback
        _emit(run_id, -1, f'Error: {e}', 'error')


@app.route('/api/progress/<run_id>')
def progress_stream(run_id):
    def generate():
        sent     = 0
        deadline = time.time() + 300
        while time.time() < deadline:
            with _lock:
                events = _runs.get(run_id, {}).get('events', [])
                done   = _runs.get(run_id, {}).get('done', False)
            while sent < len(events):
                yield f"data: {json.dumps(events[sent])}\n\n"
                sent += 1
            if done:
                break
            time.sleep(0.15)
    return Response(generate(), mimetype='text/event-stream',
                    headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'})


# ════════════════════════════════════════
# Calificaciones
# ════════════════════════════════════════

@app.route('/api/rate', methods=['POST'])
def rate():
    data    = request.json
    run_id  = data.get('run_id')
    rating  = int(data.get('rating', 0))
    comment = data.get('comment', '').strip()
    history = load_history()
    for rec in history:
        if rec['run_id'] == run_id:
            rec['rating']  = rating
            rec['comment'] = comment
            break
    save_history(history)
    rebuild_excel()
    return jsonify({'ok': True})


# ════════════════════════════════════════
# Historial y métricas
# ════════════════════════════════════════

@app.route('/api/history')
def get_history():
    return jsonify(load_history())


@app.route('/api/metrics')
def get_metrics():
    history = load_history()
    rated   = [r for r in history if r.get('rating')]
    total   = len(history)
    avg_r   = round(sum(r['rating'] for r in rated) / len(rated), 2) if rated else 0
    avg_t   = round(sum(r.get('elapsed', 0) for r in history) / total, 1) if total else 0
    pct_cal = round(len(rated) / total * 100) if total else 0

    by_model = defaultdict(list)
    for r in rated:
        by_model[r['modelo']].append(r['rating'])
    q_by_model = {m: round(sum(v)/len(v), 2) for m, v in by_model.items()}

    by_date = defaultdict(int)
    for r in history:
        by_date[r['timestamp'][:10]] += 1

    trend = [{'date': r['timestamp'][:10], 'rating': r['rating'],
               'evaluado': r['evaluado']} for r in rated[-30:]]

    by_eval = defaultdict(list)
    for r in rated:
        by_eval[r['evaluado']].append(r['rating'])
    avg_by_eval = sorted(
        [{'name': e, 'avg': round(sum(v)/len(v), 2), 'count': len(v)} for e, v in by_eval.items()],
        key=lambda x: -x['avg']
    )

    last5 = [r['rating'] for r in rated[-5:]]
    alert = len(last5) >= 3 and (sum(last5)/len(last5)) < 3

    return jsonify({
        'total':         total,
        'total_rated':   len(rated),
        'avg_rating':    avg_r,
        'avg_time':      avg_t,
        'pct_calificado': pct_cal,
        'quality_by_model': q_by_model,
        'by_date':       dict(sorted(by_date.items())[-30:]),
        'trend':         trend,
        'avg_by_eval':   avg_by_eval,
        'alert':         alert,
        'times':         [r.get('elapsed', 0) for r in history[-50:] if r.get('elapsed', 0) > 0],
    })


@app.route('/api/export')
def export():
    rebuild_excel()
    return send_file(QUALITY_XL, as_attachment=True, download_name='CALIDAD DE INFORMES.xlsx')


@app.route('/download/<path:filename>')
def download(filename):
    fp = os.path.join(BASE_DIR, filename)
    if os.path.exists(fp) and filename.lower().endswith('.docx'):
        return send_file(fp, as_attachment=True)
    return jsonify({'error': 'No encontrado'}), 404


# ════════════════════════════════════════
# Arranque
# ════════════════════════════════════════

if __name__ == '__main__':
    import webbrowser
    print("\n  CCHH — Generador de Informes de Liderazgo")
    print(f"  Directorio: {BASE_DIR}")
    print("  Navegador:  http://localhost:5000\n")
    threading.Timer(1.2, lambda: webbrowser.open('http://localhost:5000')).start()
    app.run(debug=False, host='0.0.0.0', port=5000, threaded=True)
